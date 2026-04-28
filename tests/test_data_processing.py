import logging
from datetime import date, datetime

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

import main
from excel.updater import append_dataframe_to_excel
from utils.converter import normalizar_y_validar_dataset


def test_normalizar_y_validar_dataset_ordena_y_transforma_importes():
    df = pd.DataFrame(
        {
            "fecha": ["03/01/2026", "01/01/2026", "invalida"],
            "importe": ["1.234,50", "100,00", "abc"],
        }
    )

    normalizado, resumen = normalizar_y_validar_dataset(
        df,
        col_fecha="fecha",
        columnas_importe=["importe"],
        logger=logging.getLogger("test"),
    )

    fechas_validas = [f.date() for f in normalizado["fecha"].dropna()]
    assert fechas_validas == [date(2026, 1, 1), date(2026, 1, 3)]
    assert resumen["fechas_invalidas"] == 1
    assert resumen["importes_invalidos"]["importe"] == 1
    assert float(normalizado.iloc[0]["importe"]) == 100.00
    assert float(normalizado.iloc[1]["importe"]) == 1234.50


def test_normalizar_y_validar_dataset_columnas_faltantes_no_rompe_flujo():
    df = pd.DataFrame({"otra_columna": [1, 2]})

    normalizado, resumen = normalizar_y_validar_dataset(
        df,
        col_fecha="fecha",
        columnas_importe=["importe_inexistente"],
        logger=logging.getLogger("test"),
    )

    assert "otra_columna" in normalizado.columns
    assert resumen["fechas_invalidas"] == 0
    assert resumen["importes_invalidos"] == {}


def test_append_dataframe_to_excel_detecta_ultima_fila_real(tmp_path):
    archivo = tmp_path / "maestro.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"
    ws["A1"] = "Fecha"
    ws["B1"] = "Importe"
    ws["A2"] = "2026-01-01"
    ws["B2"] = 10
    ws["C5"] = "ruido"
    wb.save(archivo)

    df_nuevo = pd.DataFrame(
        [
            ["2026-01-02", 20],
            ["2026-01-03", 30],
        ],
        columns=["Fecha", "Importe"],
    )

    resultado = append_dataframe_to_excel(
        maestro_path=str(archivo),
        hoja_destino="Datos",
        df=df_nuevo,
        columna_inicio=1,
    )

    assert resultado["fila_inicio"] == 3
    assert resultado["filas_insertadas"] == 2

    wb_check = load_workbook(archivo)
    ws_check = wb_check["Datos"]
    assert ws_check["A3"].value == "2026-01-02"
    assert ws_check["B4"].value == 30


def test_append_dataframe_to_excel_hoja_inexistente_lanza_error(tmp_path):
    archivo = tmp_path / "maestro.xlsx"
    wb = Workbook()
    wb.save(archivo)

    with pytest.raises(ValueError):
        append_dataframe_to_excel(
            maestro_path=str(archivo),
            hoja_destino="NoExiste",
            df=pd.DataFrame([[1]], columns=["A"]),
            columna_inicio=1,
        )


def test_procesar_dataset_rango_invalido_devuelve_none(monkeypatch):
    class FixedDateTime(datetime):
        @classmethod
        def today(cls):
            return cls(2026, 1, 10)

    monkeypatch.setattr(main, "datetime", FixedDateTime)
    monkeypatch.setattr(main, "obtener_ultima_fecha", lambda *args, **kwargs: datetime(2026, 1, 20))

    dataset = {
        "nombre": "FACTURADOS",
        "hoja_destino": "Datos",
        "columna_fecha": "fecha",
    }

    resultado = main.procesar_dataset(
        maestro_path="dummy.xlsx",
        dataset_config=dataset,
        fecha_inicial_config=date(2025, 1, 1),
        logger=logging.getLogger("test"),
    )

    assert resultado is None


def test_procesar_dataset_permite_un_dia_pendiente(monkeypatch):
    class FixedDateTime(datetime):
        @classmethod
        def today(cls):
            return cls(2026, 4, 27)

    monkeypatch.setattr(main, "datetime", FixedDateTime)
    monkeypatch.setattr(main, "obtener_ultima_fecha", lambda *args, **kwargs: datetime(2026, 4, 25))

    dataset = {
        "nombre": "FACTURADOS",
        "hoja_destino": "Datos",
        "columna_fecha": "fecha",
    }

    resultado = main.procesar_dataset(
        maestro_path="dummy.xlsx",
        dataset_config=dataset,
        fecha_inicial_config=date(2025, 1, 1),
        logger=logging.getLogger("test"),
    )

    assert resultado == (date(2026, 4, 26), date(2026, 4, 26))


def test_main_proceso_usa_maestro_path_por_usuario(monkeypatch):
    config = {
        "web": {
            "login_url": "https://example.com/login",
            "logout_url": "https://example.com/logout",
        },
        "paths": {
            "download_dir": "downloads_temp",
            "procesados_dir": "procesados",
        },
        "procesamiento": {
            "fecha_inicial_si_vacio": "2025-01-01",
        },
        "usuarios": [
            {
                "nombre": "USR1",
                "username": "u1",
                "password": "p1",
                "maestro_path": "C:/maestros/usr1.xlsx",
            },
            {
                "nombre": "USR2",
                "username": "u2",
                "password": "p2",
                "maestro_path": "C:/maestros/usr2.xlsx",
            },
        ],
        "datasets": [
            {
                "nombre": "FACTURADOS",
                "hoja_destino": "facturadosok",
                "columna_inicio": 1,
                "columna_fecha": "Fecha",
                "columnas_importe": ["Total"],
                "tiene_formulas": False,
                "url": "https://example.com/facturados",
            }
        ],
    }

    class DummyDriver:
        def quit(self):
            return None

    estado = {"usuario_actual": None}
    maestro_paths_procesar = []
    maestro_paths_append = []

    monkeypatch.setattr(main, "cargar_config", lambda: config)
    monkeypatch.setattr(main, "setup_logger", lambda: logging.getLogger("test"))
    monkeypatch.setattr(main, "iniciar_driver", lambda *args, **kwargs: DummyDriver())

    def fake_login(driver, login_url, username, password, max_retries):
        for usr in config["usuarios"]:
            if usr["username"] == username:
                estado["usuario_actual"] = usr["nombre"]
                assert max_retries == 3
                return
        raise AssertionError("Usuario inesperado en login")

    monkeypatch.setattr(main, "login", fake_login)
    monkeypatch.setattr(main, "logout", lambda *args, **kwargs: None)

    def fake_procesar_dataset(maestro_path, dataset_config, fecha_inicial_config, logger):
        maestro_paths_procesar.append(maestro_path)
        return date(2026, 1, 1), date(2026, 1, 2)

    monkeypatch.setattr(main, "procesar_dataset", fake_procesar_dataset)

    def fake_descargar(*args, **kwargs):
        df = pd.DataFrame([["01/01/2026", "100,00"]], columns=["Fecha", "Total"])
        return estado["usuario_actual"], df

    monkeypatch.setattr(main, "descargar_y_leer_dataset", fake_descargar)
    monkeypatch.setattr(main, "normalizar_y_validar_dataset", lambda df, *args, **kwargs: (df, {}))

    def fake_append(maestro_path, hoja_destino, df, columna_inicio):
        maestro_paths_append.append(maestro_path)
        return {"fila_inicio": 2, "filas_insertadas": len(df)}

    monkeypatch.setattr(main, "append_dataframe_to_excel", fake_append)

    salida = main.main_proceso(mode="manual")

    assert salida == config["paths"]["download_dir"]
    assert maestro_paths_procesar == ["C:/maestros/usr1.xlsx", "C:/maestros/usr2.xlsx"]
    assert maestro_paths_append == ["C:/maestros/usr1.xlsx", "C:/maestros/usr2.xlsx"]
