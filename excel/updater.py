from copy import copy

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator


class MasterExcelArrayFormulaError(Exception):
    pass


def ensure_sheet_exists_with_headers(
    maestro_path,
    hoja_destino,
    headers,
    columna_inicio,
):
    wb = load_workbook(maestro_path)

    if hoja_destino in wb.sheetnames:
        return False

    ws = wb.create_sheet(title=hoja_destino)

    for offset, header in enumerate(headers, start=0):
        ws.cell(row=1, column=columna_inicio + offset, value=header)

    wb.save(maestro_path)

    return True


def append_dataframe_to_excel(
    maestro_path,
    hoja_destino,
    df,
    columna_inicio
):

    wb = load_workbook(maestro_path)

    if hoja_destino not in wb.sheetnames:
        raise ValueError(f"La hoja '{hoja_destino}' no existe en el maestro.")

    ws = wb[hoja_destino]

    # Buscar ultima fila real basada en la columna_inicio
    ultima_fila = ws.max_row
    while (
        ultima_fila > 1
        and ws.cell(row=ultima_fila, column=columna_inicio).value is None
    ):
        ultima_fila -= 1

    fila_inicio_insercion = ultima_fila + 1
    fila_base = ultima_fila if ultima_fila > 1 else None

    # Insertar datos
    for i, row in enumerate(df.itertuples(index=False), start=0):
        for j, value in enumerate(row, start=0):

            nueva_celda = ws.cell(
                row=fila_inicio_insercion + i,
                column=columna_inicio + j,
                value=value
            )

            # Copiar formato desde la ultima fila de datos si existe
            celda_base = None
            if fila_base is not None:
                celda_base = ws.cell(
                    row=fila_base,
                    column=columna_inicio + j
                )

            if celda_base is not None and celda_base.has_style:
                nueva_celda.font = copy(celda_base.font)
                nueva_celda.border = copy(celda_base.border)
                nueva_celda.fill = copy(celda_base.fill)
                nueva_celda.number_format = copy(celda_base.number_format)
                nueva_celda.protection = copy(celda_base.protection)
                nueva_celda.alignment = copy(celda_base.alignment)

    wb.save(maestro_path)

    return {
        "fila_inicio": fila_inicio_insercion,
        "filas_insertadas": len(df)
    }


def estirar_formulas(
    maestro_path,
    hoja_destino,
    fila_inicio,
    filas_insertadas
):

    wb = load_workbook(maestro_path)
    ws = wb[hoja_destino]

    fila_base = fila_inicio - 1

    for col in range(1, ws.max_column + 1):

        celda_base = ws.cell(row=fila_base, column=col)

        if celda_base.data_type == "f":

            formula_base = celda_base.value
            coordenada_base = celda_base.coordinate

            for i in range(filas_insertadas):

                nueva_fila = fila_inicio + i
                nueva_celda = ws.cell(row=nueva_fila, column=col)

                # Traducir formula ajustando referencias
                try:
                    nueva_formula = Translator(
                        formula_base,
                        origin=coordenada_base
                    ).translate_formula(
                        ws.cell(row=nueva_fila, column=col).coordinate
                    )
                except TypeError as exc:
                    if "ArrayFormula" in str(exc):
                        raise MasterExcelArrayFormulaError(
                            "No se pueden estirar formulas tipo array en el Excel maestro. "
                            "Revise la ultima fila de la hoja destino y asegure que no tenga formulas array "
                            "(entre llaves {}) en las columnas donde se insertan datos."
                        ) from exc
                    raise

                nueva_celda.value = nueva_formula

                # Copiar formato completo
                if celda_base.has_style:
                    nueva_celda.font = copy(celda_base.font)
                    nueva_celda.border = copy(celda_base.border)
                    nueva_celda.fill = copy(celda_base.fill)
                    nueva_celda.number_format = copy(celda_base.number_format)
                    nueva_celda.protection = copy(celda_base.protection)
                    nueva_celda.alignment = copy(celda_base.alignment)

    wb.save(maestro_path)
